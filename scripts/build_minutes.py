# -*- coding: utf-8 -*-
"""
video2minutes · Step 2: 结构化 JSON -> 两份 Word 文档

用法:
    python build_minutes.py meeting_info.json --ts _transcribe/transcript_ts.txt -o 输出目录

输入 meeting_info.json 描述两份文档:
    transcript_doc : 音频转文字稿 (说明 + 正文 = 转写时间戳文本 + 误识对照表)
    minutes_doc    : 会议纪要    (概况/背景/要点/待确认/下一步...)

块类型 (blocks):
    h1 / h2      标题
    p            段落
    kv           两列键值表 (会议名称/时间/地点...)
    table        普通表格 {headers, rows, aligns?}  aligns: "center"|null
    num          编号列表
    note         带边框的说明段
    sign         右对齐落款
    space        空行

见 docs/meeting_info.example.json。
"""
import argparse
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def set_font(run, size=12, bold=False, name="宋体"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def para(doc, text, size=12, bold=False, align=None, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    return p


def render_blocks(doc, blocks):
    for b in blocks:
        t = b["type"]
        if t == "h1":
            para(doc, b["text"], size=15, bold=True,
                 align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10)
        elif t == "h2":
            para(doc, b["text"], size=13, bold=True,
                 align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)
        elif t == "p":
            para(doc, b["text"])
        elif t == "kv":
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for k, v in b["items"]:
                cells = table.add_row().cells
                cells[0].width = Cm(3.5)
                cells[1].width = Cm(12)
                set_font(cells[0].paragraphs[0].add_run(k), bold=True)
                set_font(cells[1].paragraphs[0].add_run(v))
        elif t == "table":
            table = doc.add_table(rows=1, cols=len(b["headers"]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            aligns = b.get("aligns") or [None] * len(b["headers"])
            for i, h in enumerate(b["headers"]):
                cell = table.rows[0].cells[i]
                p = cell.paragraphs[0]
                if aligns[i] == "center":
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(p.add_run(h), bold=True)
            for row in b["rows"]:
                cells = table.add_row().cells
                for i, v in enumerate(row):
                    p = cells[i].paragraphs[0]
                    if aligns[i] == "center":
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_font(p.add_run(str(v)))
        elif t == "num":
            for i, text in enumerate(b["items"], 1):
                para(doc, f"{i}. {text}")
        elif t == "note":
            p = para(doc, b["text"], size=10.5)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement(f"w:{edge}")
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "4")
                e.set(qn("w:space"), "4")
                pbdr.append(e)
            pPr.append(pbdr)
        elif t == "sign":
            para(doc, b["text"], align=WD_ALIGN_PARAGRAPH.RIGHT)
        elif t == "space":
            para(doc, "")


def build_transcript_doc(info, ts_lines, out_path):
    d = info["transcript_doc"]
    doc = Document()
    para(doc, d["title"], size=18, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in d["meta"]:
        cells = table.add_row().cells
        set_font(cells[0].paragraphs[0].add_run(k), bold=True)
        set_font(cells[1].paragraphs[0].add_run(v))
    doc.add_paragraph()

    for note in d.get("notes", []):
        para(doc, note, size=10.5)

    para(doc, d["body_title"], size=15, bold=True, space_after=10)
    for line in ts_lines:
        para(doc, line, size=10.5, space_after=4)

    if d.get("fixes"):
        doc.add_page_break()
        para(doc, d["fixes_title"], size=15, bold=True, space_after=10)
        para(doc, d["fixes_note"], size=10.5)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(d["fixes_headers"]):
            set_font(table.rows[0].cells[i].paragraphs[0].add_run(h), bold=True)
        for wrong, right in d["fixes"]:
            cells = table.add_row().cells
            set_font(cells[0].paragraphs[0].add_run(wrong))
            set_font(cells[1].paragraphs[0].add_run(right))

    doc.save(out_path)


def build_minutes_doc(info, out_path):
    m = info["minutes_doc"]
    doc = Document()
    for line in m["title_lines"]:
        para(doc, line, size=18, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    if m.get("subtitle"):
        para(doc, m["subtitle"], size=12,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
    render_blocks(doc, m["blocks"])
    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="结构化 JSON -> 转写稿 + 会议纪要 docx")
    ap.add_argument("info", help="meeting_info.json 路径")
    ap.add_argument("--ts", help="transcript_ts.txt 路径 (转写稿正文)")
    ap.add_argument("-o", "--outdir", default=".", help="输出目录")
    args = ap.parse_args()

    with open(args.info, encoding="utf-8") as f:
        info = json.load(f)
    prefix = info["meeting"]["file_prefix"]
    os.makedirs(args.outdir, exist_ok=True)

    ts_lines = []
    if args.ts and os.path.exists(args.ts):
        with open(args.ts, encoding="utf-8") as f:
            ts_lines = [ln.rstrip("\n") for ln in f if ln.strip()]

    if "transcript_doc" in info:
        p = os.path.join(args.outdir, f"{prefix}-音频转文字.docx")
        build_transcript_doc(info, ts_lines, p)
        print(f"已生成: {p}")
    if "minutes_doc" in info:
        p = os.path.join(args.outdir, f"{prefix}-会议纪要.docx")
        build_minutes_doc(info, p)
        print(f"已生成: {p}")


if __name__ == "__main__":
    main()
